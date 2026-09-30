from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

def add_slide(title_text, subtitle_text, bullet_points):
    slide = prs.slides.add_slide(blank_layout)
    
    # Title
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(1.2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(32)
    p.font.color.rgb = RGBColor(0, 51, 102)
    
    # Subtitle
    if subtitle_text:
        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.size = Pt(18)
        p2.font.color.rgb = RGBColor(100, 100, 100)
        
    # Bullets
    bodyBox = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(4.5))
    tf2 = bodyBox.text_frame
    tf2.word_wrap = True
    for idx, bp in enumerate(bullet_points):
        p_bp = tf2.add_paragraph() if idx > 0 else tf2.paragraphs[0]
        p_bp.text = f"• {bp}"
        p_bp.font.size = Pt(20)
        p_bp.font.color.rgb = RGBColor(30, 30, 30)
        p_bp.space_after = Pt(14)

# Slide 1: Title
add_slide(
    "Isomorph AI: Zero-Trust Code Forensics Engine",
    "On-Device Isomorphic Syntax Tree Matching for Windows on Snapdragon",
    [
        "Enterprise-grade code security and obfuscation detection.",
        "100% On-Device, offline processing powered by Qualcomm AI Hub.",
        "Air-gapped execution avoiding IP exposure and cloud leaks."
    ]
)

# Slide 2: Problem
add_slide(
    "The Problem: Vulnerability & IP Exposure",
    "Flaws in legacy code plagiarism & comparison tools",
    [
        "Vulnerability to Obfuscation: String and token matching fail against variable renaming and AST mutations.",
        "Data Privacy Risks: Uploading code repositories to cloud LLMs creates IP leak risks.",
        "High Latency: Remote server checks add execution delay."
    ]
)

# Slide 3: Solution
add_slide(
    "Technological Solution: Quad-Factor Isomorphic Analysis",
    "Structural parsing + ONNX runtime acceleration",
    [
        "Layer 1 (35%): AST Graph Topology Matching using normalized node structures.",
        "Layer 2 (25%): Canonical IR Sequence Alignment via Longest Common Subsequence (LCS).",
        "Layer 3 (20%): Structural N-Gram TF-IDF Vectorization for localized pattern overlap.",
        "Layer 4 (20%): ONNX NPU Neural Embeddings running on QNNExecutionProvider."
    ]
)

# Slide 4: Telemetry & Performance
add_slide(
    "Real-Time Telemetry & Performance Benchmarks",
    "Tested on Snapdragon Hardware",
    [
        "Inference Speed: ~2.08 ms execution latency per code pair on Hexagon Tensor Processor.",
        "Hardware Backend: Qualcomm Neural Network Execution Provider (QNNExecutionProvider) with DirectML fallback.",
        "Privacy Context: 100% Zero-Trust Air-Gapped execution (0 external network requests)."
    ]
)

# Slide 5: Features
add_slide(
    "Enterprise Features & Multi-File Forensics",
    "End-to-End Audit Capabilities",
    [
        "Repository Zip Inspector: Automatic pairwise N x N similarity matrix calculation.",
        "Compliance Reporting: One-click PDF audit report generation detailing risk levels.",
        "Multi-Language Normalizer: Supports Python, C, C++, Java, and JavaScript."
    ]
)

prs.save("D:\\ACPDF-AI\\Isomorph_AI_Pitch_Deck.pptx")
print("Pitch deck generated successfully at D:\\ACPDF-AI\\Isomorph_AI_Pitch_Deck.pptx")
